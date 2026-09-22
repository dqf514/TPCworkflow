# -*- coding: utf-8 -*-
"""财务凭证导出映射方案：模板列、CRUD、与导出引擎衔接。"""
from __future__ import annotations

import json
import os
from copy import deepcopy
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from app.extensions import db
from app.models import FinanceExportMapping, WorkflowDefinition
from app.services.finance_export_profiles import (
    FINANCE_EXPORT_PROFILES,
    get_finance_export_profile,
    list_finance_export_profiles,
)
from app.services.form_schema import resolve_definition_schema
from app.services.form_value_format import collect_export_columns

TEMPLATE_REL_PATH = os.path.join('templates', 'finance', 'yonsuite_voucher_import.xlsx')
SHEET_NAME = '凭证查询卡片主表'
FIELD_KEY_ROW = 4
LABEL_ROW = 6

LINE_MODE_LABELS = {
    'expense_detail': '报销明细（借方按明细 + 贷方汇总）',
    'payment_single': '付款单（一借一贷）',
}

INSTANCE_FIELD_SOURCES = [
    {'key': 'display_no', 'label': '流程编号'},
    {'key': 'definition_name', 'label': '流程名称'},
    {'key': 'initiator_name', 'label': '申请人'},
    {'key': 'department_name', 'label': '申请部门'},
    {'key': 'started_at', 'label': '提交时间'},
    {'key': 'finished_at', 'label': '完成时间'},
    {'key': 'status', 'label': '流程状态'},
]

SOURCE_TYPE_LABELS = {
    'export_option': '导出参数',
    'form': '表单字段',
    'instance': '流程属性',
    'constant': '固定值',
}


def _template_abs_path() -> str:
    base = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
    return os.path.join(base, TEMPLATE_REL_PATH)


def _parse_json(raw: str | None, default: Any) -> Any:
    if not raw:
        return default
    try:
        return json.loads(raw)
    except (TypeError, json.JSONDecodeError):
        return default


def _dump_json(value: Any, default: str) -> str:
    if value is None:
        return default
    return json.dumps(value, ensure_ascii=False)


def load_template_columns() -> list[dict]:
    path = _template_abs_path()
    if not os.path.isfile(path):
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]
    seen: set[str] = set()
    out: list[dict] = []
    desc_count = 0
    for col in range(1, ws.max_column + 1):
        key = ws.cell(FIELD_KEY_ROW, col).value
        if not key:
            continue
        key = str(key).strip()
        label = str(ws.cell(LABEL_ROW, col).value or key).strip()
        uid = key
        if key == 'description':
            desc_count += 1
            uid = f'description_{desc_count}'
        if uid in seen:
            continue
        seen.add(uid)
        out.append({'key': key, 'uid': uid, 'label': label})
    wb.close()
    return out


def _line_mode_to_mode(line_mode: str | None) -> str:
    if line_mode == 'payment_single':
        return 'payment_single'
    return 'expense_detail'


def _mode_to_line_mode(mode: str | None) -> str:
    if mode == 'payment_single':
        return 'payment_single'
    return 'expense_detail'


def mapping_row_to_dict(row: FinanceExportMapping) -> dict:
    definition_ids = _parse_json(row.definition_ids, [])
    column_mappings = _parse_json(row.column_mappings, [])
    export_options = _parse_json(row.export_options, {})
    return {
        'id': row.id,
        'sector_id': row.sector_id,
        'name': row.name,
        'description': row.description or '',
        'definition_ids': definition_ids,
        'line_mode': _line_mode_to_mode(row.line_mode),
        'detail_key': row.detail_key or '',
        'amount_source_key': row.amount_source_key or 'amount',
        'column_mappings': column_mappings,
        'export_options': export_options,
        'sort_order': row.sort_order or 0,
        'is_active': bool(row.is_active),
        'created_at': row.created_at.isoformat(sep=' ', timespec='seconds') if row.created_at else None,
        'updated_at': row.updated_at.isoformat(sep=' ', timespec='seconds') if row.updated_at else None,
    }


def list_mappings_for_sector(sector_id: int) -> list[dict]:
    rows = (
        FinanceExportMapping.query.filter_by(sector_id=sector_id)
        .order_by(FinanceExportMapping.sort_order, FinanceExportMapping.id)
        .all()
    )
    return [mapping_row_to_dict(r) for r in rows]


def get_mapping(mapping_id: int, sector_id: int | None = None) -> FinanceExportMapping | None:
    q = FinanceExportMapping.query.filter_by(id=mapping_id)
    if sector_id is not None:
        q = q.filter_by(sector_id=sector_id)
    return q.first()


def create_mapping(sector_id: int, payload: dict, user_id: int | None) -> dict:
    row = FinanceExportMapping(
        sector_id=sector_id,
        name=str(payload.get('name') or '').strip() or '未命名方案',
        description=(payload.get('description') or '').strip() or None,
        definition_ids=_dump_json(payload.get('definition_ids') or [], '[]'),
        line_mode=_mode_to_line_mode(payload.get('line_mode')),
        detail_key=(payload.get('detail_key') or '').strip() or None,
        amount_source_key=(payload.get('amount_source_key') or 'amount').strip() or 'amount',
        column_mappings=_dump_json(payload.get('column_mappings') or [], '[]'),
        export_options=_dump_json(payload.get('export_options') or {}, '{}'),
        sort_order=int(payload.get('sort_order') or 0),
        is_active=bool(payload.get('is_active', True)),
        created_by=user_id,
    )
    db.session.add(row)
    db.session.commit()
    return mapping_row_to_dict(row)


def update_mapping(row: FinanceExportMapping, payload: dict) -> dict:
    if 'name' in payload:
        row.name = str(payload.get('name') or '').strip() or row.name
    if 'description' in payload:
        row.description = (payload.get('description') or '').strip() or None
    if 'definition_ids' in payload:
        row.definition_ids = _dump_json(payload.get('definition_ids') or [], '[]')
    if 'line_mode' in payload:
        row.line_mode = _mode_to_line_mode(payload.get('line_mode'))
    if 'detail_key' in payload:
        row.detail_key = (payload.get('detail_key') or '').strip() or None
    if 'amount_source_key' in payload:
        row.amount_source_key = (payload.get('amount_source_key') or 'amount').strip() or 'amount'
    if 'column_mappings' in payload:
        row.column_mappings = _dump_json(payload.get('column_mappings') or [], '[]')
    if 'export_options' in payload:
        row.export_options = _dump_json(payload.get('export_options') or {}, '{}')
    if 'sort_order' in payload:
        row.sort_order = int(payload.get('sort_order') or 0)
    if 'is_active' in payload:
        row.is_active = bool(payload.get('is_active'))
    row.updated_at = datetime.now()
    db.session.commit()
    return mapping_row_to_dict(row)


def delete_mapping(row: FinanceExportMapping) -> None:
    db.session.delete(row)
    db.session.commit()


def list_sector_workflow_definitions(sector_id: int) -> list[dict]:
    rows = (
        WorkflowDefinition.query.filter_by(business_sector_id=sector_id)
        .order_by(WorkflowDefinition.name, WorkflowDefinition.id)
        .all()
    )
    return [
        {
            'id': wf.id,
            'key': wf.key,
            'name': wf.name,
            'status': wf.status,
        }
        for wf in rows
    ]


def list_workflow_field_sources(definition_ids: list[int]) -> list[dict]:
    if not definition_ids:
        return []
    fields: list[dict] = []
    seen: set[str] = set()

    def add_field(item: dict) -> None:
        uid = f"{item.get('scope')}:{item.get('detail_key') or ''}:{item.get('key')}"
        if uid in seen:
            return
        seen.add(uid)
        fields.append(item)

    for def_id in definition_ids:
        wf = WorkflowDefinition.query.get(def_id)
        if not wf:
            continue
        schema = resolve_definition_schema(wf)
        keys, labels, field_defs = collect_export_columns(schema)
        for key in keys:
            fdef = field_defs.get(key) or {}
            ftype = fdef.get('type')
            if ftype == 'detail_total':
                src = (fdef.get('source_field') or {})
                detail_key = src.get('key') or ''
                add_field({
                    'key': key,
                    'label': labels.get(key, key),
                    'scope': 'form',
                    'detail_key': '',
                    'group': wf.name,
                })
                cols = (src.get('props') or {}).get('columns') or []
                for col in cols:
                    ckey = col.get('key')
                    if not ckey:
                        continue
                    add_field({
                        'key': ckey,
                        'label': f"{src.get('label') or detail_key}.{col.get('label') or ckey}",
                        'scope': 'detail',
                        'detail_key': detail_key,
                        'group': wf.name,
                    })
                continue
            add_field({
                'key': key,
                'label': labels.get(key, key),
                'scope': 'form',
                'detail_key': '',
                'group': wf.name,
            })

    for inst in INSTANCE_FIELD_SOURCES:
        add_field({
            'key': inst['key'],
            'label': inst['label'],
            'scope': 'instance',
            'detail_key': '',
            'group': '流程属性',
        })
    return fields


def _builtin_profile_as_mapping_template(profile_id: str) -> dict | None:
    base = get_finance_export_profile(profile_id)
    if not base:
        return None
    column_mappings: list[dict] = []
    template_cols = {c['key']: c for c in load_template_columns()}
    for opt in base.get('options') or []:
        key = opt.get('key')
        if not key or opt.get('type') == 'switch':
            continue
        col = template_cols.get(key)
        column_mappings.append({
            'template_key': key,
            'template_label': col['label'] if col else key,
            'source_type': 'export_option',
            'source_key': key,
            'constant_value': '',
            'lookup': {},
        })
    expense_map = base.get('expense_account_map') or {}
    if expense_map:
        col = template_cols.get('accSubject_code')
        column_mappings.append({
            'template_key': 'accSubject_code',
            'template_label': col['label'] if col else 'accSubject_code',
            'source_type': 'form',
            'source_key': 'expense_category',
            'constant_value': '',
            'lookup': expense_map,
        })
    export_options = {}
    for opt in base.get('options') or []:
        k = opt.get('key')
        if k and opt.get('default') is not None:
            export_options[k] = opt.get('default')
    return {
        'name': f"{base.get('name')}（自定义）",
        'description': base.get('description') or '',
        'line_mode': base.get('mode') or 'expense_detail',
        'detail_key': 'expense_items',
        'amount_source_key': 'amount',
        'definition_ids': [],
        'column_mappings': column_mappings,
        'export_options': export_options,
    }


def mapping_to_export_profile(row: FinanceExportMapping) -> dict:
    export_options = _parse_json(row.export_options, {})
    base_id = export_options.pop('_base_profile_id', None)
    base = get_finance_export_profile(base_id) if base_id else None
    if not base:
        for p in FINANCE_EXPORT_PROFILES:
            if _mode_to_line_mode(p.get('mode')) == (row.line_mode or 'expense_detail'):
                base = p
                break
    profile = deepcopy(base) if base else {
        'id': f'custom:{row.id}',
        'name': row.name,
        'mode': _line_mode_to_mode(row.line_mode),
        'options': [],
    }
    profile['id'] = f'custom:{row.id}'
    profile['name'] = row.name
    profile['description'] = row.description or profile.get('description', '')
    profile['definition_ids'] = _parse_json(row.definition_ids, [])
    profile['mode'] = _line_mode_to_mode(row.line_mode)
    profile['detail_key'] = row.detail_key or 'expense_items'
    profile['amount_source_key'] = row.amount_source_key or 'amount'
    profile['column_mappings'] = _parse_json(row.column_mappings, [])
    profile['custom'] = True

    expense_map: dict[str, str] = {}
    for cm in profile['column_mappings']:
        lookup = cm.get('lookup') or {}
        if not isinstance(lookup, dict):
            continue
        tk = cm.get('template_key')
        if tk == 'accSubject_code' and lookup:
            expense_map.update({str(k): str(v) for k, v in lookup.items()})
    if expense_map:
        profile['expense_account_map'] = expense_map

    for opt in profile.get('options') or []:
        k = opt.get('key')
        if k in export_options:
            opt['default'] = export_options[k]
    profile['export_option_defaults'] = export_options
    return profile


def resolve_finance_export_profile(profile_id: str) -> dict | None:
    if profile_id.startswith('custom:'):
        try:
            mid = int(profile_id.split(':', 1)[1])
        except (TypeError, ValueError):
            return None
        row = FinanceExportMapping.query.filter_by(id=mid, is_active=True).first()
        if not row:
            return None
        return mapping_to_export_profile(row)
    return get_finance_export_profile(profile_id)


def list_profiles_for_sector(sector_id: int | None) -> list[dict]:
    out = list_finance_export_profiles()
    for item in out:
        item['source'] = 'builtin'
    if not sector_id:
        return out
    custom_rows = (
        FinanceExportMapping.query.filter_by(sector_id=sector_id, is_active=True)
        .order_by(FinanceExportMapping.sort_order, FinanceExportMapping.id)
        .all()
    )
    for row in custom_rows:
        profile = mapping_to_export_profile(row)
        out.append({
            'id': profile['id'],
            'name': profile['name'],
            'description': profile.get('description', ''),
            'definition_keys': profile.get('definition_keys') or [],
            'definition_names': profile.get('definition_names') or [],
            'definition_ids': profile.get('definition_ids') or [],
            'mode': profile.get('mode'),
            'options': deepcopy(profile.get('options') or []),
            'source': 'custom',
            'mapping_id': row.id,
        })
    return out


def finance_mapping_meta() -> dict:
    return {
        'line_modes': [
            {'value': 'expense_detail', 'label': LINE_MODE_LABELS['expense_detail']},
            {'value': 'payment_single', 'label': LINE_MODE_LABELS['payment_single']},
        ],
        'source_types': [
            {'value': k, 'label': v} for k, v in SOURCE_TYPE_LABELS.items()
        ],
        'builtin_templates': [
            {'id': p['id'], 'name': p['name']} for p in FINANCE_EXPORT_PROFILES
        ],
    }
