# -*- coding: utf-8 -*-
"""生成用友 YonSuite BIP 凭证导入 Excel。"""
from __future__ import annotations

import io
import os
import re
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from app.models import Department, WorkflowInstance
from app.services.finance_export_profiles import get_finance_export_profile, instance_matches_profile
from app.services.finance_export_mapping_service import resolve_finance_export_profile
from app.services.form_templates.expense_reimbursement import EXPENSE_CATEGORY_OPTIONS
from app.services.form_value_format import extract_dept_ids, money_numeric
from app.services.graph_utils import parse_variables

TEMPLATE_REL_PATH = os.path.join('templates', 'finance', 'yonsuite_voucher_import.xlsx')
SHEET_NAME = '凭证查询卡片主表'
FIELD_KEY_ROW = 4  # 1-based
DATA_START_ROW = 10  # 1-based，模板示例数据从第 10 行起
CURRENCY_VALUE = '人民币##CNY'


def _template_abs_path() -> str:
    base = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
    return os.path.join(base, TEMPLATE_REL_PATH)


def _load_column_map(ws) -> tuple[dict[str, int], int | None, int | None]:
    mapping: dict[str, int] = {}
    header_desc_col: int | None = None
    line_desc_col: int | None = None
    desc_seen = 0
    for col in range(1, ws.max_column + 1):
        key = ws.cell(FIELD_KEY_ROW, col).value
        if not key:
            continue
        key = str(key).strip()
        if key == 'description':
            desc_seen += 1
            if desc_seen == 1:
                header_desc_col = col
            else:
                line_desc_col = col
            continue
        mapping[key] = col
    if line_desc_col is None:
        line_desc_col = header_desc_col
    return mapping, header_desc_col, line_desc_col


def _set_cell(row_cells: dict[int, Any], col_map: dict[str, int], field: str, value: Any) -> None:
    if value is None or value == '':
        return
    col = col_map.get(field)
    if col:
        row_cells[col] = value


def _period_from_dt(dt: datetime | None) -> str:
    if not dt:
        dt = datetime.now()
    return dt.strftime('%Y-%m')


def _make_time_from_inst(inst: WorkflowInstance) -> datetime:
    dt = inst.finished_at or inst.started_at or getattr(inst, 'updated_at', None) or datetime.now()
    if isinstance(dt, datetime):
        return dt.replace(hour=0, minute=0, second=0, microsecond=0)
    return datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)


def _resolve_dept(inst: WorkflowInstance, form: dict, dept_cache: dict[int, Department | None]) -> tuple[str, str]:
    raw = form.get('cost_department') or form.get('request_department') or inst.department_id
    dept_ids = extract_dept_ids(raw)
    if not dept_ids and inst.department_id:
        dept_ids = [inst.department_id]
    if not dept_ids:
        return '', ''
    dept_id = dept_ids[0]
    if dept_id not in dept_cache:
        dept_cache[dept_id] = Department.query.get(dept_id)
    dept = dept_cache.get(dept_id)
    name = dept.name if dept else str(dept_id)
    code = str(dept_id)
    return code, name


def _option_label(field: dict, raw: Any) -> str:
    if raw is None or raw == '':
        return ''
    options = (field.get('props') or {}).get('options') or []
    raw_str = str(raw)
    for opt in options:
        code = str(opt.get('code') if opt.get('code') is not None else opt.get('value') or '')
        if code == raw_str:
            return str(opt.get('name') or opt.get('label') or code)
    return raw_str


def _voucher_header(
    voucher_id: str,
    inst: WorkflowInstance,
    form: dict,
    opts: dict[str, Any],
    make_time: datetime,
) -> dict[str, Any]:
    reason = str(form.get('reason') or inst.definition.name if inst.definition else '流程凭证').strip()
    display_no = inst.display_no or str(inst.id)
    return {
        'id': voucher_id,
        'org_code': opts.get('org_code', ''),
        'accBook_code': opts.get('accBook_code', ''),
        'period': opts.get('period') or _period_from_dt(make_time),
        'makeTime': make_time,
        'voucherType_voucherstr': opts.get('voucherType_voucherstr') or '记',
        'billCode': opts.get('billCode') or display_no[-3:] if display_no else '',
        'description': reason[:200],
    }


def _amount_line(
    header: dict[str, Any],
    col_map: dict[str, int],
    header_desc_col: int | None,
    line_desc_col: int | None,
    *,
    acc_subject: str,
    line_desc: str,
    debit: float = 0,
    credit: float = 0,
    dept_code: str = '',
    dept_name: str = '',
    bank_account: str = '',
) -> dict[int, Any]:
    cells: dict[int, Any] = {}
    for k, v in header.items():
        if k == 'description':
            if header_desc_col:
                cells[header_desc_col] = v
            continue
        _set_cell(cells, col_map, k, v)
    _set_cell(cells, col_map, 'accSubject_code', acc_subject)
    if line_desc_col:
        cells[line_desc_col] = line_desc[:255]
    _set_cell(cells, col_map, 'currency_code', CURRENCY_VALUE)
    if debit > 0:
        _set_cell(cells, col_map, 'debitOrg', round(debit, 2))
        _set_cell(cells, col_map, 'debitOriginal', round(debit, 2))
        _set_cell(cells, col_map, 'creditOrg', 0)
        _set_cell(cells, col_map, 'creditOriginal', 0)
    elif credit > 0:
        _set_cell(cells, col_map, 'debitOrg', 0)
        _set_cell(cells, col_map, 'debitOriginal', 0)
        _set_cell(cells, col_map, 'creditOrg', round(credit, 2))
        _set_cell(cells, col_map, 'creditOriginal', round(credit, 2))
    if dept_code or dept_name:
        _set_cell(cells, col_map, '0001', dept_code)
        _set_cell(cells, col_map, '0001_auxName', dept_name)
    if bank_account:
        _set_cell(cells, col_map, 'bankAccount', bank_account)
    return cells


def _lines_for_expense(
    inst: WorkflowInstance,
    form: dict,
    profile: dict,
    opts: dict[str, Any],
    col_map: dict[str, int],
    header_desc_col: int | None,
    line_desc_col: int | None,
    dept_cache: dict[int, Department | None],
) -> list[dict[int, Any]]:
    voucher_id = f"WF{inst.display_no or inst.id}"
    make_time = _make_time_from_inst(inst)
    header = _voucher_header(voucher_id, inst, form, opts, make_time)
    dept_code, dept_name = _resolve_dept(inst, form, dept_cache)
    if opts.get('dept_aux_code'):
        dept_code = str(opts['dept_aux_code'])
    elif dept_name and not dept_code:
        dept_code = dept_name

    account_map = profile.get('expense_account_map') or {}
    default_debit = str(opts.get('debit_accSubject_code') or '6602')
    credit_subject = str(opts.get('credit_accSubject_code') or '224102')
    bank_subject = str(opts.get('bank_accSubject_code') or '100201')

    payee = str(form.get('payee_name') or form.get('payee') or '').strip()
    bank_account = str(form.get('bank_account') or '').strip()
    reason = str(form.get('reason') or '').strip()

    detail_key = profile.get('detail_key') or 'expense_items'
    amount_key = profile.get('amount_source_key') or 'amount'
    rows = form.get(detail_key) or []
    if not isinstance(rows, list):
        rows = []

    lines: list[dict[int, Any]] = []
    total = 0.0

    type_field = {'props': {'options': EXPENSE_CATEGORY_OPTIONS}}

    for row in rows:
        if not isinstance(row, dict):
            continue
        amount = money_numeric(row.get(amount_key) or row.get('amount'))
        if not amount or amount <= 0:
            continue
        total += amount
        expense_type = row.get('expense_type')
        acc = account_map.get(str(expense_type), default_debit)
        type_label = _option_label(type_field, expense_type) if expense_type else ''
        remark = str(row.get('remark') or '').strip()
        line_desc = ' - '.join(x for x in [type_label, remark, reason] if x) or reason or '报销'
        lines.append(
            _amount_line(
                header,
                col_map,
                header_desc_col,
                line_desc_col,
                acc_subject=acc,
                line_desc=line_desc,
                debit=amount,
                dept_code=dept_code,
                dept_name=dept_name,
            ),
        )

    if total <= 0:
        amount = money_numeric(form.get('total_amount') or form.get('payment_amount'))
        if amount and amount > 0:
            total = amount
            lines.append(
                _amount_line(
                    header,
                    col_map,
                    header_desc_col,
                    line_desc_col,
                    acc_subject=default_debit,
                    line_desc=reason or '报销',
                    debit=total,
                    dept_code=dept_code,
                    dept_name=dept_name,
                ),
            )

    if total <= 0:
        return []

    credit_desc = f"报销-{payee}" if payee else f"报销-{reason or voucher_id}"
    lines.append(
        _amount_line(
            header,
            col_map,
            header_desc_col,
            line_desc_col,
            acc_subject=credit_subject,
            line_desc=credit_desc[:255],
            credit=total,
            dept_code=dept_code,
            dept_name=dept_name,
        ),
    )

    if opts.get('include_bank_line') and bank_account:
        lines.append(
            _amount_line(
                header,
                col_map,
                header_desc_col,
                line_desc_col,
                acc_subject=credit_subject,
                line_desc=f"冲应付-{payee or reason}"[:255],
                debit=total,
            ),
        )
        lines.append(
            _amount_line(
                header,
                col_map,
                header_desc_col,
                line_desc_col,
                acc_subject=bank_subject,
                line_desc=f"付款-{payee or bank_account}"[:255],
                credit=total,
                bank_account=bank_account,
            ),
        )

    return lines


def _lines_for_payment(
    inst: WorkflowInstance,
    form: dict,
    opts: dict[str, Any],
    col_map: dict[str, int],
    header_desc_col: int | None,
    line_desc_col: int | None,
    dept_cache: dict[int, Department | None],
) -> list[dict[int, Any]]:
    amount = money_numeric(form.get('payment_amount'))
    if not amount or amount <= 0:
        return []

    voucher_id = f"WF{inst.display_no or inst.id}"
    make_time = _make_time_from_inst(inst)
    header = _voucher_header(voucher_id, inst, form, opts, make_time)
    dept_code, dept_name = _resolve_dept(inst, form, dept_cache)
    if opts.get('dept_aux_code'):
        dept_code = str(opts['dept_aux_code'])

    payee = str(form.get('payee') or '').strip()
    reason = str(form.get('reason') or '').strip()
    bank_account = str(form.get('bank_account') or '').strip()

    debit_subject = str(opts.get('debit_accSubject_code') or '122105')
    credit_subject = str(opts.get('credit_accSubject_code') or '224102')

    desc_debit = f"付款-{payee}-{reason}" if payee else reason or '付款'
    desc_credit = desc_debit

    lines = [
        _amount_line(
            header,
            col_map,
            header_desc_col,
            line_desc_col,
            acc_subject=debit_subject,
            line_desc=desc_debit[:255],
            debit=amount,
            dept_code=dept_code,
            dept_name=dept_name,
        ),
        _amount_line(
            header,
            col_map,
            header_desc_col,
            line_desc_col,
            acc_subject=credit_subject,
            line_desc=desc_credit[:255],
            credit=amount,
            dept_code=dept_code,
            dept_name=dept_name,
            bank_account=bank_account,
        ),
    ]
    return lines


def _clear_sample_data(ws, start_row: int = DATA_START_ROW) -> None:
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)


def _write_rows(ws, start_row: int, row_cells_list: list[dict[int, Any]]) -> None:
    row_idx = start_row
    for cells in row_cells_list:
        for col, val in cells.items():
            ws.cell(row_idx, col, val)
        row_idx += 1


def build_yonsuite_voucher_workbook(
    instances: list[WorkflowInstance],
    profile_id: str,
    runtime_options: dict[str, Any],
) -> bytes:
    profile = resolve_finance_export_profile(profile_id)
    if not profile:
        profile = get_finance_export_profile(profile_id)
    if not profile:
        raise ValueError('未知的财务导出方案')

    template_path = _template_abs_path()
    if not os.path.isfile(template_path):
        raise FileNotFoundError('未找到用友凭证导入模板，请联系管理员')

    wb = load_workbook(template_path)
    ws = wb[SHEET_NAME]
    col_map, header_desc_col, line_desc_col = _load_column_map(ws)

    dept_cache: dict[int, Department | None] = {}
    all_lines: list[dict[int, Any]] = []
    approved_only = runtime_options.get('approved_only', True)
    if isinstance(approved_only, str):
        approved_only = approved_only.lower() in ('1', 'true', 'yes', 'on')

    for inst in instances:
        if not instance_matches_profile(inst, profile):
            continue
        if approved_only and inst.status != 'approved':
            continue
        variables = parse_variables(inst.variables)
        form = variables.get('form') or {}
        mode = profile.get('mode')
        if mode == 'expense_detail':
            all_lines.extend(
                _lines_for_expense(
                    inst, form, profile, runtime_options, col_map,
                    header_desc_col, line_desc_col, dept_cache,
                ),
            )
        elif mode == 'payment_single':
            all_lines.extend(
                _lines_for_payment(
                    inst, form, runtime_options, col_map,
                    header_desc_col, line_desc_col, dept_cache,
                ),
            )

    if not all_lines:
        raise ValueError('没有符合导出条件的流程数据，请检查流程类型与状态筛选')

    _clear_sample_data(ws)
    _write_rows(ws, DATA_START_ROW, all_lines)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sanitize_export_filename(name: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', '_', name or '凭证导入')
    return cleaned[:80]
